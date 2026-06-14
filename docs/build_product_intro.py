from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "日报表OCR产品介绍.docx"
ASSETS = ROOT / "docs" / "assets"

BLUE = RGBColor(46, 85, 180)
NAVY = RGBColor(15, 23, 42)
MUTED = RGBColor(87, 99, 120)
LIGHT_BLUE = "EEF4FF"
LIGHT_GRAY = "F7F9FC"
BORDER = "D9E2EF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, color=None, bold=None, name="Microsoft YaHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def add_paragraph(doc, text="", size=10.5, color=NAVY, bold=False, after=6, before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        size = 16
        color = BLUE
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        size = 13
        color = BLUE
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        size = 11.5
        color = NAVY
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    for run in p.runs:
        set_run_font(run, size=10.5, color=NAVY)
    if not p.runs:
        run = p.add_run(text)
    else:
        p.runs[0].text = text
        run = p.runs[0]
    set_run_font(run, size=10.5, color=NAVY)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=NAVY)
    return p


def add_callout(doc, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.3)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color="C7D8F5")
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=NAVY, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    set_run_font(run, size=9, color=MUTED)


def add_picture(doc, filename, caption, width=6.25):
    image_path = ASSETS / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    add_caption(doc, caption)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_GRAY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_run_font(run, size=9.5, color=NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_run_font(run, size=9.3, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
normal.font.size = Pt(10.5)

header = section.header.paragraphs[0]
header.text = "日报表 OCR 产品介绍"
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_run_font(header.runs[0], size=9, color=MUTED)
footer = section.footer.paragraphs[0]
footer.text = "面向供应商的日报表数字化工具"
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run_font(footer.runs[0], size=9, color=MUTED)

add_paragraph(doc, "客户介绍资料", size=10.5, color=BLUE, bold=True, after=4)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(6)
run = p.add_run("日报表 OCR")
set_run_font(run, size=28, color=NAVY, bold=True)
add_paragraph(doc, "把手写日报表变成可校对、可入库、可导出的电子数据", size=14, color=MUTED, after=18)

meta = doc.add_table(rows=2, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.autofit = False
for col in meta.columns:
    col.width = Inches(3.15)
for row in meta.rows:
    for cell in row.cells:
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_border(cell, "C7D8F5")
        set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
for cell, text in zip(meta.rows[0].cells + meta.rows[1].cells, [
    "适用对象：供应商、生产班组、日报录入人员",
    "核心价值：少手抄、少错录、快入库",
    "使用方式：网页端上传，手机端可拍照上传",
    "输出结果：结构化数据与 Excel 文件",
]):
    run = cell.paragraphs[0].add_run(text)
    set_run_font(run, size=10, color=NAVY, bold=True)

add_callout(doc, "一句话说明：供应商只需要上传或拍照提交日报表，系统自动识别内容，人工确认后即可形成可管理的数据。")

add_heading(doc, "1. 产品定位", 1)
add_paragraph(doc, "日报表 OCR 是一套面向供应商日报表场景的识别和数据管理工具。它不要求使用人员了解 IT 技术，也不改变原有填写习惯。供应商仍然可以使用纸质日报表、手写表格或拍照文件，系统负责把这些材料转成后续可查询、可导出的电子数据。")
add_paragraph(doc, "适合的日常场景包括：生产日报、派工单、供应商交付日报、班组填报记录、来料或加工数量登记等。")

add_heading(doc, "2. 核心功能", 1)
add_table(
    doc,
    ["功能", "客户看到的效果", "带来的价值"],
    [
        ["图片 / PDF 上传", "支持电脑选择文件，手机端可直接拍照上传。", "减少扫描、传文件、人工整理的步骤。"],
        ["批量识别", "一次可以上传多张日报表，系统逐张识别。", "适合每天多张表、多个班组一起处理。"],
        ["人工校对", "左侧看原图，右侧改数据，像 Excel 一样核对。", "保留人工确认环节，避免误入库。"],
        ["模板字段管理", "按日报表样式配置表头和明细字段。", "不同供应商或不同表格可以统一口径。"],
        ["数据管理", "所有上传记录、状态和入库数据集中查看。", "方便追溯、筛选、导出。"],
        ["Excel 导出", "确认后的数据可导出为 Excel。", "便于对账、统计和内部流转。"],
    ],
    [1.35, 2.45, 2.45],
)

add_heading(doc, "3. 使用流程", 1)
add_number(doc, "上传日报表：电脑端选择图片或 PDF；手机端可选择图片，也可以直接拍照上传。")
add_number(doc, "开始识别：点击“开始识别”，系统自动提交 OCR 和结构化处理。")
add_number(doc, "校对结果：在校对页左侧查看原图，右侧修改识别后的字段和明细。")
add_number(doc, "提交入库：确认无误后点击“提交入库”，数据进入管理列表。")
add_number(doc, "查询与导出：在数据管理页面查看历史记录，并按需要导出 Excel。")

add_picture(doc, "01-ocr-upload-desktop.png", "图 1：电脑端 OCR 上传页面，选择报表文件后即可开始识别。", width=6.25)
add_picture(doc, "04-mobile-camera-upload.png", "图 2：手机端可直接点击“拍照上传”，适合现场提交日报表。", width=2.25)

add_heading(doc, "4. 校对方式更适合一线人员", 1)
add_paragraph(doc, "日报表数据最终要进入管理系统，因此系统保留了人工校对环节。校对页面采用左右对照方式：左边显示原始图片，右边显示识别后的数据。用户可以一边看原图，一边修改字段。")
add_bullet(doc, "原图和数据放在同一个页面，不需要来回切换窗口。")
add_bullet(doc, "中间分隔条可以拖动，方便根据图片大小调整左右区域。")
add_bullet(doc, "明细区采用紧凑表格，操作方式接近 Excel，降低学习成本。")
add_bullet(doc, "发现识别不准时，直接在单元格里修改，再提交入库。")
add_picture(doc, "03-review-workbench-desktop.png", "图 3：校对页面，左侧原图、右侧识别结果，便于人工核对。", width=6.25)

add_heading(doc, "5. 数据管理与导出", 1)
add_paragraph(doc, "所有上传过的报表都会保留处理状态，例如“待校对”“已入库”“失败”等。管理人员可以查看详情、删除无用记录，也可以将已确认的数据导出为 Excel。")
add_picture(doc, "02-data-management-desktop.png", "图 4：数据管理页面，可查看上传记录、状态，并导出 Excel。", width=6.25)

add_heading(doc, "6. 给供应商的使用建议", 1)
add_table(
    doc,
    ["建议", "说明"],
    [
        ["拍照清晰", "尽量让表格完整入镜，避免强反光、遮挡和严重倾斜。"],
        ["先校对再入库", "识别结果需要人工确认，尤其是数量、日期和产品名称。"],
        ["按模板填写", "表头、产品名称、数量等字段尽量写在固定位置，识别效果会更稳定。"],
        ["定期导出", "可以按日或按周导出 Excel，用于对账和统计。"],
    ],
    [1.55, 4.7],
)

add_heading(doc, "7. 产品价值总结", 1)
add_callout(doc, "对供应商来说，它不是一套复杂系统，而是一个“拍照 / 上传 - 校对 - 入库 - 导出”的简单工具。")
add_bullet(doc, "减少人工录入：把重复抄表、打字录入的工作交给系统。")
add_bullet(doc, "降低出错风险：识别后仍由人工确认，关键数据不直接自动入库。")
add_bullet(doc, "提高流转效率：日报表从纸面材料变成电子数据，后续查询和导出更方便。")
add_bullet(doc, "便于统一管理：多张报表、多名录入人员、多种状态集中在一个页面。")

doc.save(OUT)
print(OUT)
