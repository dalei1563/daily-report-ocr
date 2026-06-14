from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "日报表OCR部署说明.docx"

BLUE = RGBColor(46, 85, 180)
NAVY = RGBColor(15, 23, 42)
MUTED = RGBColor(87, 99, 120)
CODE_BG = "F4F6F9"
HEADER_BG = "EEF4FF"
BORDER = "D9E2EF"


def set_run_font(run, size=None, color=None, bold=None, name="Microsoft YaHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def border_cell(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def margins_cell(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def p(doc, text="", size=10.5, color=NAVY, bold=False, after=6, before=0):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.15
    run = para.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    return para


def h(doc, text, level=1):
    para = doc.add_paragraph()
    if level == 1:
        size, before, after, color = 16, 16, 8, BLUE
    elif level == 2:
        size, before, after, color = 13, 12, 6, BLUE
    else:
        size, before, after, color = 11.5, 8, 4, NAVY
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    run = para.add_run(text)
    set_run_font(run, size=size, color=color, bold=True)
    return para


def bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.12
    run = para.add_run(text)
    set_run_font(run, size=10.3, color=NAVY)


def number(doc, text):
    para = doc.add_paragraph(style="List Number")
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.12
    run = para.add_run(text)
    set_run_font(run, size=10.3, color=NAVY)


def code(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.3)
    cell = table.cell(0, 0)
    shade_cell(cell, CODE_BG)
    border_cell(cell, "D7DBE2")
    margins_cell(cell, top=100, bottom=100, start=140, end=140)
    lines = text.strip("\n").split("\n")
    cell.paragraphs[0].text = ""
    for i, line in enumerate(lines):
        para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(line)
        set_run_font(run, size=9, color=NAVY, name="Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def table(doc, headers, rows, widths):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for i, width in enumerate(widths):
        tbl.columns[i].width = Inches(width)
    for i, text in enumerate(headers):
        cell = tbl.cell(0, i)
        shade_cell(cell, HEADER_BG)
        border_cell(cell)
        margins_cell(cell)
        run = cell.paragraphs[0].add_run(text)
        set_run_font(run, size=9.5, color=NAVY, bold=True)
    for row in rows:
        cells = tbl.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            border_cell(cell)
            margins_cell(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            run = cell.paragraphs[0].add_run(text)
            set_run_font(run, size=9.2, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
normal.font.size = Pt(10.5)

header = section.header.paragraphs[0]
header.text = "日报表 OCR 部署说明"
set_run_font(header.runs[0], size=9, color=MUTED)
footer = section.footer.paragraphs[0]
footer.text = "建议部署前由实施人员核对域名、端口、密钥和备份策略"
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run_font(footer.runs[0], size=9, color=MUTED)

p(doc, "实施交付文档", size=10.5, color=BLUE, bold=True, after=4)
title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(6)
r = title.add_run("日报表 OCR 部署说明")
set_run_font(r, size=26, color=NAVY, bold=True)
p(doc, "适用于本地服务器、局域网服务器或容器化环境部署", size=13, color=MUTED, after=16)

table(
    doc,
    ["项目", "说明"],
    [
        ["系统组成", "前端 Web 页面 + 后端 FastAPI 服务 + SQLite 数据库 + 本地文件目录"],
        ["推荐方式", "Docker Compose 部署，前端 8080 端口，后端 8000 端口"],
        ["默认账号", "admin / admin123，正式使用前必须修改或停用默认账号"],
        ["外部依赖", "PaddleOCR 服务密钥、大模型 API Key"],
    ],
    [1.45, 4.85],
)

h(doc, "1. 部署前准备", 1)
bullet(doc, "服务器建议使用 Windows Server 或 Linux 服务器，需能访问 PaddleOCR 与大模型服务地址。")
bullet(doc, "推荐安装 Docker 与 Docker Compose；如不使用 Docker，可采用 Python + Node.js 手动部署。")
bullet(doc, "准备 PaddleOCR Token、大模型 API Key、生产环境 APP_SECRET_KEY。")
bullet(doc, "确认开放端口：前端 8080，后端 8000；如通过域名访问，可由 Nginx 或网关统一转发。")

h(doc, "2. 目录与数据说明", 1)
p(doc, "系统运行时会保存数据库、上传原图、OCR 原始结果和导出文件。Docker 部署时，仓库根目录下的 data 文件夹会挂载到容器内 /data。")
table(
    doc,
    ["目录 / 文件", "用途", "是否建议备份"],
    [
        ["data/app.db", "SQLite 数据库，保存账号、模板、上传记录、识别结果。", "必须备份"],
        ["data/uploads", "用户上传的图片或 PDF 原文件。", "必须备份"],
        ["data/ocr_raw", "OCR 原始返回文件和中间结果。", "建议备份"],
        ["data/exports", "导出的 Excel 文件。", "可按需备份"],
        [".env", "部署环境变量和密钥。", "需安全保存，不要外泄"],
    ],
    [1.5, 3.15, 1.65],
)

h(doc, "3. 推荐部署方式：Docker Compose", 1)
number(doc, "复制环境变量模板。")
code(doc, "Copy-Item .env.example .env")
number(doc, "编辑 .env 文件，至少填写以下内容。")
code(doc, """PADDLE_OCR_TOKEN=你的PaddleOCR密钥
LLM_API_KEY=你的大模型密钥
APP_SECRET_KEY=请替换为随机长字符串
DATABASE_URL=sqlite:////data/app.db
DATA_DIR=/data
BACKEND_CORS_ORIGINS=http://服务器IP:8080""")
number(doc, "构建并启动服务。")
code(doc, "docker compose up --build -d")
number(doc, "查看服务状态。")
code(doc, "docker compose ps\ndocker compose logs -f backend\ndocker compose logs -f frontend")
p(doc, "启动成功后，浏览器访问前端地址：http://服务器IP:8080。后端接口文档地址：http://服务器IP:8000/docs。")

h(doc, "4. 手动部署方式：不用 Docker", 1)
h(doc, "4.1 后端服务", 2)
code(doc, """cd backend
python -m venv .venv
.\\.venv\\Scripts\\activate
pip install -r requirements.txt
$env:DATABASE_URL="sqlite:///./app.db"
$env:DATA_DIR="./data"
$env:APP_SECRET_KEY="请替换为随机长字符串"
uvicorn app.main:app --host 0.0.0.0 --port 8000""")
h(doc, "4.2 前端服务", 2)
code(doc, """cd frontend
npm install
npm run build
npm run preview -- --host 0.0.0.0 --port 8080""")
p(doc, "手动部署适合测试或临时演示。正式生产环境建议使用 Docker Compose 或由专业 Web 服务器托管前端 dist 目录，并将 /api 转发到后端 8000 端口。")

h(doc, "5. 首次登录与系统配置", 1)
number(doc, "使用默认账号登录：admin / admin123。")
number(doc, "进入管理员页面，配置 PaddleOCR 与大模型密钥。")
number(doc, "创建正式账号并绑定模板。")
number(doc, "正式使用前，修改默认管理员密码，或新增管理员后停用默认账号。")
number(doc, "上传一张测试日报表，完成识别、校对、入库和 Excel 导出，确认闭环可用。")

h(doc, "6. 环境变量说明", 1)
table(
    doc,
    ["变量名", "示例", "说明"],
    [
        ["PADDLE_OCR_TOKEN", "******", "PaddleOCR 服务密钥。"],
        ["PADDLE_OCR_JOB_URL", "https://paddleocr.aistudio-app.com/api/v2/ocr/jobs", "PaddleOCR 任务提交地址。"],
        ["PADDLE_OCR_MODEL", "PaddleOCR-VL-1.6", "OCR 模型名称。"],
        ["LLM_BASE_URL", "https://api.deepseek.com", "兼容 OpenAI 协议的大模型服务地址。"],
        ["LLM_API_KEY", "******", "大模型调用密钥。"],
        ["LLM_MODEL", "deepseek-v4-flash", "大模型名称。"],
        ["APP_SECRET_KEY", "随机长字符串", "登录令牌签名密钥，生产环境必须更换。"],
        ["DATABASE_URL", "sqlite:////data/app.db", "数据库连接地址。"],
        ["DATA_DIR", "/data", "上传、OCR、导出文件保存目录。"],
        ["BACKEND_CORS_ORIGINS", "http://服务器IP:8080", "允许访问后端的前端地址。"],
    ],
    [1.55, 2.35, 2.4],
)

h(doc, "7. 启动验证清单", 1)
bullet(doc, "打开 http://服务器IP:8080 能看到登录页面。")
bullet(doc, "使用管理员账号可以登录。")
bullet(doc, "管理员页面中 OCR 与大模型配置显示已配置。")
bullet(doc, "上传图片后可以开始识别。")
bullet(doc, "识别完成后进入校对页面，左侧能看到原图，右侧能编辑数据。")
bullet(doc, "提交入库后，数据管理页面能看到已入库记录。")
bullet(doc, "点击导出 Excel 能下载文件。")

h(doc, "8. 备份与升级建议", 1)
bullet(doc, "每日备份 data/app.db 和 data/uploads；如存储空间允许，也备份 data/ocr_raw。")
bullet(doc, "升级前先停止服务，并完整备份 data 目录和 .env 文件。")
bullet(doc, "升级代码后重新执行 docker compose up --build -d。")
bullet(doc, "升级后用一张测试报表验证上传、识别、校对、入库、导出流程。")

h(doc, "9. 常见问题", 1)
table(
    doc,
    ["问题", "可能原因", "处理方式"],
    [
        ["无法登录", "默认账号被修改、数据库未初始化或服务未连接。", "检查后端日志、数据库文件和账号状态。"],
        ["提示缺少识别密钥", "PaddleOCR 或大模型 Key 未配置。", "进入管理员页面配置密钥，或检查 .env。"],
        ["识别失败", "外部 OCR 队列繁忙、密钥错误或网络不可达。", "查看记录错误信息和 backend 日志，稍后重试或检查网络。"],
        ["图片不显示", "后端文件接口异常或上传文件丢失。", "检查 data/uploads 是否存在对应文件。"],
        ["导出失败", "data/exports 无写入权限或后端异常。", "检查数据目录权限和后端日志。"],
        ["手机无法拍照上传", "浏览器权限或非 HTTPS/局域网限制。", "使用手机浏览器访问部署地址，允许相机权限；必要时配置 HTTPS。"],
    ],
    [1.55, 2.2, 2.55],
)

h(doc, "10. 生产环境安全提醒", 1)
bullet(doc, "不要把真实密钥写入代码仓库或截图中。")
bullet(doc, "APP_SECRET_KEY 必须更换为随机长字符串。")
bullet(doc, "正式上线后不要继续使用默认密码 admin123。")
bullet(doc, "如部署在公网，建议使用 HTTPS，并限制管理后台访问范围。")
bullet(doc, "定期备份 data 目录，确认备份文件可恢复。")

doc.save(OUT)
print(OUT)
